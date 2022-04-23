from openpyxl import Workbook
from openpyxl.chart import LineChart, Series, Reference
# result = [1000, 900, 800, 700, 600]
# all_speed = [200, 250, 300, 400, 450]


def excell_table(result, all_speed):

    wb = Workbook(write_only=True)
    ws = wb.create_sheet()

    rows = []
    rows.append(('Number of events', 'Speed'))
    for i in range(0, len(result), 1):
        rows.append((result[i], all_speed[i]))

    for row in rows:
        ws.append(row)

    chart1 = LineChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "Results of test"
    chart1.y_axis.title = 'Number of events'
    chart1.x_axis.title = 'Speed of traffic'

    data = Reference(ws, min_col=1, min_row=1, max_row=len(result)+1)
    cats = Reference(ws, min_col=2, min_row=2, max_row=len(result)+1, max_col=2)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    chart1.shape = 4
    ws.add_chart(chart1, "K10")

    wb.save("perf_chart.xlsx")

    print("**********FINISH************")
    print("table configured")

# excell_table(result, all_speed)
